"""End-to-end vertical slice test (PRD 21):
create deal -> import CSV -> approve mappings -> edit assumptions -> calculate
DCF -> approve model version -> export XLSX.
"""

import io

CSV_CONTENT = b"""fiscal_year,statement_type,label,value,currency,unit_scale
2024,income_statement,Total revenue,10000000,USD,actuals
2024,income_statement,Cost of revenue,4000000,USD,actuals
2024,income_statement,Sales and marketing,2500000,USD,actuals
2024,income_statement,Research and development,1500000,USD,actuals
2024,income_statement,Depreciation and amortization,300000,USD,actuals
2024,income_statement,Income tax,400000,USD,actuals
2024,balance_sheet,Cash and cash equivalents,5000000,USD,actuals
2024,balance_sheet,Accounts receivable,800000,USD,actuals
2024,balance_sheet,Accounts payable,500000,USD,actuals
2024,cash_flow,Capital expenditures,250000,USD,actuals
"""


def _create_workspace_and_deal(client, headers):
    ws = client.post("/api/workspaces", json={"name": "Acme Capital"}, headers=headers).json()
    deal = client.post(
        "/api/deals",
        json={
            "workspace_id": ws["id"],
            "company_name": "TestCo SaaS",
            "industry": "saas",
            "currency": "USD",
            "unit_scale": "actuals",
            "valuation_date": "2025-01-01",
            "fiscal_year_end": "12-31",
            "forecast_years": 3,
        },
        headers=headers,
    ).json()
    return ws, deal


def _default_assumption_years(base_fy=2024, n=3, growth=0.20, margin=0.60):
    payload = []
    for i in range(1, n + 1):
        fy = base_fy + i
        for key, val in [
            ("revenue_growth_rate", growth),
            ("gross_margin", margin),
            ("opex_pct_of_revenue", 0.35),
            ("d_and_a_pct_of_revenue", 0.03),
            ("capex_pct_of_revenue", 0.02),
            ("nwc_pct_of_revenue", 0.05),
            ("tax_rate", 0.25),
        ]:
            payload.append({"key": key, "fiscal_year": fy, "value": val, "origin": "user"})
    return payload


def test_full_vertical_slice(auth_client):
    client, headers = auth_client
    ws, deal = _create_workspace_and_deal(client, headers)
    deal_id = deal["id"]
    assert deal["ready_for_valuation"] is True

    # 1. Upload a source document.
    doc_resp = client.post(
        f"/api/deals/{deal_id}/documents",
        params={"document_type": "financial_statements"},
        files={"file": ("financials.csv", io.BytesIO(CSV_CONTENT), "text/csv")},
        headers=headers,
    )
    assert doc_resp.status_code == 201
    document_id = doc_resp.json()["id"]

    # 2. Import financials from the same CSV.
    import_resp = client.post(
        f"/api/deals/{deal_id}/financials/import",
        params={"document_id": document_id},
        files={"file": ("financials.csv", io.BytesIO(CSV_CONTENT), "text/csv")},
        headers=headers,
    )
    assert import_resp.status_code == 201, import_resp.text
    import_result = import_resp.json()
    assert import_result["periods_created"] == 1
    assert import_result["line_items_created"] == 10
    assert import_result["unresolved_mappings"] == 0, import_result["warnings"]

    # 3. Review mapping, approve all.
    periods = client.get(f"/api/deals/{deal_id}/financials", headers=headers).json()
    assert len(periods) == 1
    line_item_ids = [li["id"] for li in periods[0]["line_items"]]
    approve_resp = client.post(
        f"/api/deals/{deal_id}/financials/approve-mappings", json=line_item_ids, headers=headers
    )
    assert approve_resp.status_code == 200
    assert all(li["mapping_status"] == "approved" for li in approve_resp.json())

    # 4. Data quality checks should now pass (no missing periods / mixed currency).
    checks = client.get(f"/api/deals/{deal_id}/financials/model-checks", headers=headers).json()
    assert not checks["has_blocking_failures"]

    # 5. Set base/bull/bear assumptions.
    for scenario, growth in [("base", 0.20), ("bull", 0.30), ("bear", 0.10)]:
        resp = client.put(
            f"/api/deals/{deal_id}/assumptions/{scenario}",
            json={
                "scenario": scenario,
                "wacc": 0.12,
                "terminal_growth_rate": 0.04,
                "assumptions": _default_assumption_years(growth=growth),
            },
            headers=headers,
        )
        assert resp.status_code == 200, resp.text

    # 6. Create a model version and calculate.
    mv = client.post(f"/api/deals/{deal_id}/model-versions", json={"name": "Initial Draft"}, headers=headers).json()
    calc_resp = client.post(f"/api/model-versions/{mv['id']}/calculate", headers=headers)
    assert calc_resp.status_code == 200, calc_resp.text
    scenarios = calc_resp.json()["scenarios"]
    assert set(scenarios.keys()) == {"base", "bull", "bear"}

    base = scenarios["base"]
    assert len(base["forecast"]) == 3
    assert base["dcf"]["enterprise_value"] > 0
    # Bull case should produce a higher enterprise value than bear, holding WACC/TGR fixed.
    assert scenarios["bull"]["dcf"]["enterprise_value"] > scenarios["bear"]["dcf"]["enterprise_value"]

    # 7. Sensitivities.
    sens = client.get(
        f"/api/model-versions/{mv['id']}/sensitivities/wacc-vs-growth", params={"scenario": "base"}, headers=headers
    )
    assert sens.status_code == 200
    grid = sens.json()
    assert len(grid["grid"]) == 5
    assert len(grid["grid"][0]) == 5

    # 8. Approve the model version.
    approve_mv = client.post(f"/api/model-versions/{mv['id']}/approve", headers=headers)
    assert approve_mv.status_code == 200
    assert approve_mv.json()["status"] == "approved"

    # 9. Export XLSX.
    export_resp = client.post(f"/api/model-versions/{mv['id']}/exports/xlsx", headers=headers)
    assert export_resp.status_code == 201, export_resp.text
    export_job_id = export_resp.json()["export_job_id"]

    download = client.get(f"/api/exports/{export_job_id}", headers=headers)
    assert download.status_code == 200
    assert download.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert len(download.content) > 1000  # a real workbook, not an empty stub

    # The exported workbook should open with openpyxl and contain the required tabs.
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(download.content))
    assert "11_DCF" in wb.sheetnames
    assert "16_Model_Checks" in wb.sheetnames
    assert "02_Assumptions" in wb.sheetnames


def test_import_rejects_unsupported_file_type(auth_client):
    client, headers = auth_client
    ws, deal = _create_workspace_and_deal(client, headers)
    resp = client.post(
        f"/api/deals/{deal['id']}/financials/import",
        files={"file": ("financials.txt", io.BytesIO(b"not a csv"), "text/plain")},
        headers=headers,
    )
    assert resp.status_code == 400


def test_cannot_approve_model_version_before_calculating(auth_client):
    client, headers = auth_client
    ws, deal = _create_workspace_and_deal(client, headers)
    mv = client.post(f"/api/deals/{deal['id']}/model-versions", json={"name": "v1"}, headers=headers).json()
    resp = client.post(f"/api/model-versions/{mv['id']}/approve", headers=headers)
    assert resp.status_code == 400


def test_terminal_growth_must_be_below_wacc(auth_client):
    client, headers = auth_client
    ws, deal = _create_workspace_and_deal(client, headers)
    resp = client.put(
        f"/api/deals/{deal['id']}/assumptions/base",
        json={"scenario": "base", "wacc": 0.05, "terminal_growth_rate": 0.05, "assumptions": []},
        headers=headers,
    )
    assert resp.status_code == 400
